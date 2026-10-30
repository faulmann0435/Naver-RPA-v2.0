"""자연이랑 (supplier portal nscm.62life.com): its 출고작업관리 Excel -> the Naver-shaped order table.

Pure functions, no I/O beyond parsing xlsx bytes. The table has the Naver export's column names, so
core.pipeline.process_orders (purchase orders) and core.invoice_match (tracking numbers) run unchanged.

    is_withnature_export(bytes) -> bool        does this xlsx look like the portal's export?
    read_withnature_export(bytes) -> DataFrame raw columns, all text
    to_order_frame(raw) -> DataFrame           Naver column names
    courier_code(name) -> str | None           the portal's 택배업체코드
    build_portal_upload(rows) -> bytes         the 일괄등록 sheet (주문번호 / 운송장번호 / 택배업체코드)

The export's 보내는사람주소 / 문의전화 columns (the portal's own office) are never read.
"""
from __future__ import annotations

import re
from collections.abc import Iterable
from io import BytesIO

import openpyxl
import pandas as pd

CHANNEL = "자연이랑"
PRODUCT_NO_PREFIX = "자연이랑-"  # keeps the portal's product codes apart from Naver product numbers
READY_STATUS = "발송대기"
MESSAGE_SEPARATOR = " / "
ZIP_WIDTH = 5

DETECT_COLUMNS = ("주문번호", "받는사람주소", "공급단가", "상품코드")
REQUIRED_COLUMNS = ("주문번호", "받는사람", "받는사람주소", "우편번호", "상품코드", "상품명", "주문수량")

# Column order of the table handed to the pipeline (the names Naver's export uses).
ORDER_COLUMNS = ("상품주문번호", "주문번호", "구매자명", "수취인명", "우편번호", "상품명", "옵션정보", "수량",
                 "수취인연락처1", "통합배송지", "배송메세지", "주문상태", "배송비 묶음번호", "상품번호",
                 "수취인연락처2")

UPLOAD_HEADER = ("주문번호", "운송장번호", "택배업체코드")

# Courier name (any spelling containing one of the keys) -> the portal's code. First hit wins.
_COURIER_KEYWORDS: tuple[tuple[str, str], ...] = (
    ("우체국", "01"), ("CJ", "02"), ("씨제이", "02"), ("대한통운", "02"), ("한진", "03"), ("로젠", "04"),
    ("롯데", "05"), ("DHL", "06"), ("FEDEX", "07"), ("UPS", "08"), ("일양", "09"), ("부릉", "10"),
    ("대신", "11"), ("경동", "12"), ("합동", "13"), ("새벽", "14"), ("천일", "15"), ("야쿠르트", "16"),
    ("HY", "16"), ("농협", "17"), ("에이치피엠", "18"), ("HPM", "18"),
)
_ASCII_KEYS = {"DHL", "FEDEX", "UPS", "HY", "HPM"}  # matched as whole words so "GROUPS" is not UPS


def _cell(value) -> str:
    """Cell -> trimmed text; whole numbers lose the '.0', blanks become ''."""
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _sheet_rows(data: bytes) -> list[tuple]:
    if not data.startswith(b"PK"):
        raise ValueError("엑셀(.xlsx) 파일이 아닙니다.")
    try:
        workbook = openpyxl.load_workbook(BytesIO(data), data_only=True, read_only=True)
    except Exception as exc:  # any unreadable-file error becomes one readable message
        raise ValueError(f"엑셀 파일을 읽지 못했습니다: {exc}") from exc
    sheet = workbook.worksheets[0]
    sheet.reset_dimensions()
    return list(sheet.iter_rows(values_only=True))


def _is_withnature_header(header: Iterable) -> bool:
    names = {_cell(c) for c in header}
    return all(column in names for column in DETECT_COLUMNS)


def is_withnature_export(xlsx_bytes: bytes) -> bool:
    """True when the first sheet's header row has 주문번호, 받는사람주소, 공급단가 and 상품코드."""
    try:
        rows = _sheet_rows(xlsx_bytes)
    except ValueError:
        return False
    return bool(rows) and _is_withnature_header(rows[0])


def _zip_text(value) -> str:
    text = _cell(value)
    return text.zfill(ZIP_WIDTH) if text.isdigit() and len(text) < ZIP_WIDTH else text


def read_withnature_export(xlsx_bytes: bytes) -> pd.DataFrame:
    """The export as it is: every column text (우편번호 keeps its leading zeros), blank rows dropped."""
    rows = _sheet_rows(xlsx_bytes)
    if not rows or not _is_withnature_header(rows[0]):
        raise ValueError("자연이랑 출고작업관리 파일이 아닙니다. '주문번호'·'받는사람주소'·'공급단가'·'상품코드' 열을 찾지 못했습니다.")
    header = [_cell(c) for c in rows[0]]
    missing = [c for c in REQUIRED_COLUMNS if c not in header]
    if missing:
        raise ValueError(f"자연이랑 파일에 필요한 열이 없습니다: {', '.join(missing)}")
    keep = [i for i, name in enumerate(header) if name]
    body = [[_cell(row[i]) if i < len(row) else "" for i in keep] for row in rows[1:]]
    body = [r for r in body if any(r)]
    frame = pd.DataFrame(body, columns=[header[i] for i in keep], dtype=object)
    frame["우편번호"] = [_zip_text(v) for v in frame["우편번호"]]
    return frame


def restore_address(text: str) -> str:
    """The portal swaps commas for '*': '(판교동* 아파트)' -> '(판교동, 아파트)'."""
    return text.replace("* ", ", ").replace("*", ",")


def join_messages(*parts: str) -> str:
    """Non-blank parts joined with ' / ', repeated ones once."""
    kept = list(dict.fromkeys(p.strip() for p in parts if p and p.strip()))
    return MESSAGE_SEPARATOR.join(kept)


def _quantity(text: str) -> int:
    try:
        return int(float(text))
    except ValueError as exc:
        raise ValueError(f"주문수량이 숫자가 아닙니다: {text!r}") from exc


def _column(raw: pd.DataFrame, name: str) -> list[str]:
    return [_cell(v) for v in raw[name]] if name in raw.columns else [""] * len(raw)


def to_order_frame(raw: pd.DataFrame) -> pd.DataFrame:
    """Raw export -> table with Naver's column names: ids/zip/phones text, 수량 int64."""
    order_no, product = _column(raw, "주문번호"), _column(raw, "상품명")
    mobile, landline = _column(raw, "휴대전화"), _column(raw, "전화번호")
    frame = pd.DataFrame({
        "상품주문번호": order_no,
        "주문번호": order_no,
        "구매자명": _column(raw, "보내는사람"),
        "수취인명": _column(raw, "받는사람"),
        "우편번호": [_zip_text(v) for v in _column(raw, "우편번호")],
        "상품명": product,
        "옵션정보": product,
        "수량": pd.Series([_quantity(v) for v in _column(raw, "주문수량")], dtype="int64"),
        "수취인연락처1": [m or p for m, p in zip(mobile, landline, strict=True)],
        "통합배송지": [restore_address(v) for v in _column(raw, "받는사람주소")],
        "배송메세지": [join_messages(a, b) for a, b in zip(_column(raw, "배송메시지"), _column(raw, "전달메모"), strict=True)],
        "주문상태": READY_STATUS,
        "배송비 묶음번호": order_no,
        "상품번호": [PRODUCT_NO_PREFIX + code for code in _column(raw, "상품코드")],
        "수취인연락처2": landline,
    }, columns=list(ORDER_COLUMNS))
    return frame.astype({c: object for c in frame.columns if c != "수량"})


def courier_code(name) -> str | None:
    """The portal's 택배업체코드 for a courier name ('한진택배', 'CJ대한통운' ...); None when unknown."""
    text = re.sub(r"\s+", "", str(name or "")).upper()
    for key, code in _COURIER_KEYWORDS:
        hit = re.search(rf"(?<![A-Z]){key}(?![A-Z])", text) if key in _ASCII_KEYS else key in text
        if hit:
            return code
    return None


def build_portal_upload(rows: Iterable[tuple[str, str, str]]) -> bytes:
    """xlsx for the portal's 일괄등록: header 주문번호 / 운송장번호 / 택배업체코드, then (주문번호, 운송장번호, 코드) as text."""
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "운송장등록"
    sheet.append(list(UPLOAD_HEADER))
    for row in rows:
        sheet.append([str(v) for v in row])
    for line in sheet.iter_rows(min_row=2):
        for cell in line:
            cell.number_format = "@"
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
