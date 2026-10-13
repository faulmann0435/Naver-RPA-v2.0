"""송장 매칭: match the courier's tracking-number file (한진 송장파일) to the Naver order rows.

Pure functions, no I/O beyond parsing xlsx/xls bytes. Streamlit and the Telegram bot share this core.
Three tracking-file layouts are read: the 한진 file, a purchase order a supplier sent back with a
한진택배 column, and a supplier's shipping list (택배사 / 운송장번호 / 이름 / 연락처) that mixes couriers.

    read_invoice_file(bytes)   -> list[Invoice]        one per box
    read_order_export(bytes)   -> DataFrame            the Naver order export (decrypt before)
    normalize_orders(df)       -> DataFrame            ids/zip/phones as text (also for the API table)
    match_invoices(df, inv)    -> MatchResult
    build_upload_workbook(...) -> bytes                the "발송처리" sheet Naver's 엑셀 일괄발송 reads
"""
from __future__ import annotations

import re
from collections import defaultdict
from dataclasses import dataclass
from datetime import datetime, timedelta, timezone
from io import BytesIO

import openpyxl
import pandas as pd
import xlrd

KST = timezone(timedelta(hours=9))  # own copy: the Streamlit app ships without worker/
CARRIER_NAME = "한진택배"
DELIVERY_METHOD_LABEL = "택배,등기,소포"
UPLOAD_SHEET = "발송처리"
ORDER_SHEET = "발주발송관리"
STATUS_MATCHED, STATUS_CHECK = "MATCHED", "NEEDS_CHECK"
NOTE_NO_ORDER = "주문을 못 찾음"
READY_STATUS = "발송대기"
HEADER_SCAN_ROWS = 10  # the supplier list has a totals block above its header

COL_ID, COL_ORDER_ID = "상품주문번호", "주문번호"
COL_METHOD, COL_CARRIER, COL_TRACKING = "배송방법", "택배사", "송장번호"
COL_NAME, COL_ZIP, COL_ADDRESS = "수취인명", "우편번호", "통합배송지"
COL_PHONE1, COL_PHONE2 = "수취인연락처1", "수취인연락처2"
COL_BUNDLE, COL_STATUS, COL_ITEM = "배송비 묶음번호", "주문상태", "상품명"
COL_OPTION = "옵션정보"
REASON_BY_ITEM = "같은 분 상자 여러 개 · 상품명으로 고름"
REQUIRED_ORDER_COLUMNS = (COL_ID, COL_NAME, COL_ZIP, COL_ADDRESS, COL_PHONE1)
TEXT_COLUMNS = (COL_ID, COL_ORDER_ID, COL_ZIP, COL_BUNDLE, "상품번호", "구매자연락처",
                COL_PHONE1, COL_PHONE2, COL_TRACKING)

HANJIN_TRACKING = "운송장번호"
HANJIN_REQUIRED = ("받으시는 분", "받는분우편번호", "받는분총주소")
# Courier names as suppliers write them -> (name Naver's 엑셀 일괄발송 expects, Commerce API code).
# Codes from Naver's delivery-company list (LOGEN was merged into KGB).
_HANJIN = (CARRIER_NAME, "HANJIN")
_CJ = ("CJ대한통운", "CJGLS")
_EPOST = ("우체국택배", "EPOST")
_LOGEN = ("로젠택배", "KGB")
_LOTTE = ("롯데택배", "HYUNDAI")
CARRIERS: dict[str, tuple[str, str]] = {
    "한진택배": _HANJIN, "한진": _HANJIN,
    "CJ대한통운": _CJ, "씨제이대한통운": _CJ, "대한통운": _CJ, "CJ택배": _CJ, "CJ": _CJ,
    "우체국택배": _EPOST, "우체국": _EPOST,
    "로젠택배": _LOGEN, "로젠": _LOGEN,
    "롯데택배": _LOTTE, "롯데": _LOTTE,
}
COURIER_CODES = {name: code for name, code in CARRIERS.values()}
# A supplier's shipping list: one row per box, several couriers mixed, no address.
LIST_TRACKING, LIST_CARRIER, LIST_NAME = HANJIN_TRACKING, "택배사", "이름"
LIST_REQUIRED = (LIST_CARRIER, LIST_NAME, "연락처")
# A purchase order the supplier sends back with the tracking numbers in a column named after the courier.
RETURNED_TRACKING, RETURNED_NAME = CARRIER_NAME, "받는사람"
RETURNED_REQUIRED = (RETURNED_NAME, "전화번호", "주소")
_BRACKET_TAIL = re.compile(r"\s*\[[^\]]*\]\s*$")


@dataclass(frozen=True)
class Invoice:
    """One box on the courier's file."""
    seq: str
    tracking: str
    name: str
    phone: str
    zip_code: str
    address: str
    item_name: str
    mobile: str = ""
    carrier: str = CARRIER_NAME  # Naver's name for the courier


@dataclass(frozen=True)
class Assignment:
    """A group of product orders (one shipment) and the tracking number(s) found for it."""
    product_order_ids: tuple[str, ...]
    tracking: str
    recipient: str
    items: str
    status: str
    includable: bool
    reason: str = ""
    carrier: str = CARRIER_NAME


@dataclass(frozen=True)
class UnshippedOrder:
    """A shipment group no tracking number was found for."""
    product_order_ids: tuple[str, ...]
    recipient: str
    items: str


@dataclass(frozen=True)
class MatchResult:
    assignments: tuple[Assignment, ...]
    invoices_without_order: tuple[Invoice, ...]
    orders_without_invoice: tuple[UnshippedOrder, ...]
    invoices_without_tracking: tuple[Invoice, ...] = ()  # rows of the file with no tracking number

    def matched(self) -> tuple[Assignment, ...]:
        return tuple(a for a in self.assignments if a.status == STATUS_MATCHED)

    def needs_check(self) -> tuple[Assignment, ...]:
        return tuple(a for a in self.assignments if a.status == STATUS_CHECK)

    def includable_checks(self) -> tuple[Assignment, ...]:
        return tuple(a for a in self.needs_check() if a.includable)

    def includable(self) -> tuple[Assignment, ...]:
        return tuple(a for a in self.assignments if a.includable)


# ---------------------------------------------------------------- normalizing

def norm_name(value) -> str:
    return re.sub(r"\s+", "", _text(value))


def norm_phone(value) -> str:
    return re.sub(r"\D", "", _text(value))


def norm_tracking(value) -> str:
    """Suppliers write 6006-6339-2321; Naver wants the bare number."""
    return re.sub(r"[^0-9A-Za-z]", "", _text(value))


def carrier_of(value) -> tuple[str, str]:
    """Supplier's courier name -> (Naver name, API code); ValueError for one we do not know."""
    key = re.sub(r"\s+", "", _text(value)).upper()
    for alias, carrier in CARRIERS.items():
        if alias.upper() == key:
            return carrier
    raise ValueError(f"모르는 택배사입니다: '{_text(value)}'. 개발자에게 알려 주세요.")


def norm_zip(value) -> str:
    digits = re.sub(r"\D", "", _text(value))
    return digits.zfill(5) if digits else ""


def norm_order_address(value) -> str:
    return re.sub(r"\s+", "", _text(value))


def norm_invoice_address(value) -> str:
    """Hanjin appends ' [동, 지번]' to the address; drop it before comparing."""
    return norm_order_address(_BRACKET_TAIL.sub("", _text(value)))


def _blank(value) -> bool:
    if value is None:
        return True
    try:
        return bool(pd.isna(value))
    except (TypeError, ValueError):
        return False


def _text(value) -> str:
    """Cell -> text; whole numbers lose the '.0', blanks become ''."""
    if _blank(value):
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def normalize_orders(df: pd.DataFrame) -> pd.DataFrame:
    """Copy of the order table: id/zip/phone columns as text (zip zero-padded), blanks as ''."""
    out = df.reset_index(drop=True).astype(object)
    for column in out.columns:
        if column in TEXT_COLUMNS:
            out[column] = [_text(v) for v in out[column]]
        else:
            out[column] = [("" if _blank(v) else v) for v in out[column]]
    if COL_ZIP in out.columns:
        out[COL_ZIP] = [norm_zip(v) for v in out[COL_ZIP]]
    return out


# ---------------------------------------------------------------- reading files

def _open_sheet(data: bytes, preferred: str | None = None):
    if not data.startswith(b"PK"):
        raise ValueError("엑셀(.xlsx) 파일이 아닙니다. .xls 파일은 .xlsx로 저장한 뒤 올려 주세요.")
    try:
        workbook = openpyxl.load_workbook(BytesIO(data), data_only=True, read_only=True)
    except Exception as exc:  # any unreadable-file error becomes one readable message
        raise ValueError(f"엑셀 파일을 읽지 못했습니다: {exc}") from exc
    sheet = workbook[preferred] if preferred in workbook.sheetnames else workbook.worksheets[0]
    sheet.reset_dimensions()  # Naver's export declares a one-cell size; read-only mode would stop there
    return sheet


def _sheet_rows(data: bytes) -> list[tuple]:
    """Rows of the first sheet of an .xlsx or an old .xls (suppliers' files are often slightly damaged)."""
    if data.startswith(b"\xd0\xcf\x11\xe0"):
        try:
            book = xlrd.open_workbook(file_contents=data, ignore_workbook_corruption=True)
        except Exception as exc:  # any unreadable-file error becomes one readable message
            raise ValueError(f"엑셀(.xls) 파일을 읽지 못했습니다: {exc}") from exc
        sheet = book.sheet_by_index(0)
        return [tuple(sheet.row_values(r)) for r in range(sheet.nrows)]
    return list(_open_sheet(data).iter_rows(values_only=True))


def _find_header(rows: list[tuple], key: str) -> int | None:
    for position, row in enumerate(rows[:HEADER_SCAN_ROWS]):
        if any(_text(cell) == key for cell in row):
            return position
    return None


def read_invoice_file(data: bytes) -> list[Invoice]:
    """The courier's xlsx, or a purchase order a supplier sent back with a 한진택배 column.

    One Invoice per non-empty row; rows without a tracking number come back with tracking ''
    so match_invoices can report them (a supplier may have left a box out).
    """
    rows = _sheet_rows(data)
    tracking_at = _find_header(rows, HANJIN_TRACKING)
    if tracking_at is not None:
        header = [_text(c) for c in rows[tracking_at]]
        if LIST_NAME in header and LIST_CARRIER in header:
            return _read_rows(rows, tracking_at, LIST_REQUIRED, _list_invoice)
        return _read_rows(rows, tracking_at, HANJIN_REQUIRED, _hanjin_invoice)
    returned_at = _find_header(rows, RETURNED_TRACKING)
    if returned_at is not None and RETURNED_NAME in [_text(c) for c in rows[returned_at]]:
        return _read_rows(rows, returned_at, RETURNED_REQUIRED, _returned_invoice)
    raise ValueError("송장 파일이 아닙니다. 한진 송장파일('운송장번호' 열) 또는 "
                     "'한진택배' 열에 송장번호를 적은 발주서를 보내 주세요.")


def _read_rows(rows: list[tuple], header_at: int, required: tuple[str, ...], make) -> list[Invoice]:
    header = [_text(c) for c in rows[header_at]]
    missing = [c for c in required if c not in header]
    if missing:
        raise ValueError(f"송장 파일에 필요한 열이 없습니다: {', '.join(missing)}")
    invoices = []
    for line_no, row in enumerate(rows[header_at + 1:], start=header_at + 2):
        cells = dict(zip(header, row, strict=False))
        if any(_text(v) for v in cells.values()):
            invoices.append(make(cells, line_no))
    return invoices


def _hanjin_invoice(cells: dict, line_no: int) -> Invoice:
    return Invoice(
        seq=_text(cells.get("순번")) or str(line_no), tracking=norm_tracking(cells.get(HANJIN_TRACKING)),
        name=_text(cells.get("받으시는 분")), phone=_text(cells.get("받으시는 분 전화")),
        mobile=_text(cells.get("받는분핸드폰")), zip_code=norm_zip(cells.get("받는분우편번호")),
        address=_text(cells.get("받는분총주소")), item_name=_text(cells.get("품목명")))


def _returned_invoice(cells: dict, line_no: int) -> Invoice:
    """A purchase-order row (받는사람/전화번호/주소/구분); it has no zip code."""
    return Invoice(
        seq=str(line_no), tracking=norm_tracking(cells.get(RETURNED_TRACKING)), name=_text(cells.get(RETURNED_NAME)),
        phone=_text(cells.get("전화번호")), zip_code="", address=_text(cells.get("주소")),
        item_name=_text(cells.get("구분")))


def _list_invoice(cells: dict, line_no: int) -> Invoice:
    """A supplier's shipping-list row (발송일자/택배사/운송장번호/이름/연락처/품목및수량): no zip, no address."""
    tracking = norm_tracking(cells.get(LIST_TRACKING))
    carrier = carrier_of(cells.get(LIST_CARRIER))[0] if tracking else CARRIER_NAME
    return Invoice(
        seq=str(line_no), tracking=tracking, name=_text(cells.get(LIST_NAME)), phone=_text(cells.get("연락처")),
        zip_code="", address="", item_name=_text(cells.get("품목및수량")), carrier=carrier)


def read_order_export(data: bytes) -> pd.DataFrame:
    """The Naver order export (already decrypted): every column, ids/zip/phones as text."""
    rows = list(_open_sheet(data, ORDER_SHEET).iter_rows(values_only=True))
    header_at = _find_header(rows, COL_ID)
    if header_at is None:
        raise ValueError("네이버 주문서가 아닙니다. '상품주문번호' 열을 찾지 못했습니다.")
    header = [_text(c) for c in rows[header_at]]
    keep = [i for i, name in enumerate(header) if name]
    body = [[(row[i] if i < len(row) else None) for i in keep] for row in rows[header_at + 1:]]
    body = [r for r in body if any(not _blank(v) for v in r)]
    return normalize_orders(pd.DataFrame(body, columns=[header[i] for i in keep]))


# ---------------------------------------------------------------- matching

@dataclass(frozen=True)
class _Group:
    ids: tuple[str, ...]
    recipient: str
    name: str
    zip_code: str
    address: str
    phones: frozenset[str]
    items: str
    item_text: str = ""  # every product name and option of the group, for telling boxes apart


def _items_summary(names: list[str]) -> str:
    first = next((n for n in names if n), "")
    return first if len(names) <= 1 else f"{first} 외 {len(names) - 1}건"


def _eligible(orders: pd.DataFrame) -> pd.DataFrame:
    mask = pd.Series(True, index=orders.index)
    if COL_STATUS in orders.columns:
        mask &= orders[COL_STATUS].map(_text) == READY_STATUS
    if COL_TRACKING in orders.columns:
        mask &= orders[COL_TRACKING].map(_text) == ""
    return orders[mask]


def _group_orders(orders: pd.DataFrame) -> list[_Group]:
    missing = [c for c in REQUIRED_ORDER_COLUMNS if c not in orders.columns]
    if missing:
        raise ValueError(f"주문서에 필요한 열이 없습니다: {', '.join(missing)}")
    buckets: dict[tuple, list[dict]] = defaultdict(list)
    for row in _eligible(orders).to_dict("records"):
        bundle = _text(row.get(COL_BUNDLE))
        key = ("B", bundle) if bundle else (
            "N", norm_name(row[COL_NAME]), norm_zip(row[COL_ZIP]), norm_order_address(row[COL_ADDRESS]))
        buckets[key].append(row)
    return [_group(rows) for rows in buckets.values()]


def _group(rows: list[dict]) -> _Group:
    first = rows[0]
    phones = {norm_phone(r.get(c)) for r in rows for c in (COL_PHONE1, COL_PHONE2)}
    return _Group(
        ids=tuple(_text(r[COL_ID]) for r in rows), recipient=_text(first[COL_NAME]),
        name=norm_name(first[COL_NAME]), zip_code=norm_zip(first[COL_ZIP]),
        address=norm_order_address(first[COL_ADDRESS]), phones=frozenset(p for p in phones if p),
        items=_items_summary([_text(r.get(COL_ITEM)) for r in rows]),
        item_text=" ".join(_text(r.get(c)) for r in rows for c in (COL_ITEM, COL_OPTION)))


def _mismatch(group: _Group, invoice: Invoice) -> tuple[bool, bool]:
    """(phone_ok, address_ok)"""
    phones = {norm_phone(invoice.phone), norm_phone(invoice.mobile)} - {""}
    address_ok = not invoice.address or norm_invoice_address(invoice.address) == group.address
    return bool(phones & group.phones), address_ok


def _is_candidate(group: _Group, invoice: Invoice) -> bool:
    """Same name and zip; a file without zip codes needs the phone or the address to agree as well."""
    if norm_name(invoice.name) != group.name:
        return False
    zip_code = norm_zip(invoice.zip_code)
    if zip_code:
        return zip_code == group.zip_code
    phone_ok, address_ok = _mismatch(group, invoice)
    return phone_ok or (bool(invoice.address) and address_ok)


def _bigrams(text: str) -> set[str]:
    squashed = re.sub(r"[\s\W_]+", "", text.lower())
    return {squashed[i:i + 2] for i in range(len(squashed) - 1)}


def _item_score(group: _Group, invoice: Invoice) -> float:
    """Share of the box's item text (the shorter one) found in the order's product names and options."""
    box = _bigrams(invoice.item_name)
    return len(box & _bigrams(group.item_text)) / len(box) if box else 0.0


def _best(scores: dict[int, float]) -> int | None:
    """The key with the strictly highest positive score, if there is one."""
    ranked = sorted(scores.items(), key=lambda kv: kv[1], reverse=True)
    if not ranked or ranked[0][1] <= 0 or (len(ranked) > 1 and ranked[1][1] == ranked[0][1]):
        return None
    return ranked[0][0]


def _pick_by_item(groups: list[_Group], candidates: list[list[int]], invoices: list[Invoice]) -> set[int]:
    """Same person, several boxes: pair each of the person's orders with one box by product name.

    Repeatedly takes a pairing that is best both ways (the box fits this order best and the order
    fits this box best), then drops that box from the person's other orders; an order left with one
    box gets it. Narrows `candidates` in place; returns the groups narrowed this way.
    """
    open_groups = {g for g, picked in enumerate(candidates) if len(picked) > 1}
    narrowed: set[int] = set()
    while open_groups:
        taken = {candidates[g][0] for g in narrowed}
        choices = {g: [i for i in candidates[g] if i not in taken] for g in open_groups}
        pair = next(((g, c[0]) for g, c in choices.items() if len(c) == 1), None) or _mutual_pair(
            groups, invoices, choices)
        if pair is None:
            break
        g, box = pair
        candidates[g] = [box]
        narrowed.add(g)
        open_groups.discard(g)
    return narrowed


def _mutual_pair(groups: list[_Group], invoices: list[Invoice],
                 choices: dict[int, list[int]]) -> tuple[int, int] | None:
    for g, boxes in choices.items():
        box = _best({i: _item_score(groups[g], invoices[i]) for i in boxes})
        if box is None:
            continue
        rivals = [h for h, other in choices.items() if box in other]
        if _best({h: _item_score(groups[h], invoices[box]) for h in rivals}) == g:
            return g, box
    return None


def _problem_reason(phone_ok: bool, addr_ok: bool) -> str:
    return "전화번호·주소 다름" if not phone_ok and not addr_ok else (
        "전화번호 다름" if not phone_ok else "주소 다름")


def _assignment(group: _Group, invoice: Invoice, status: str, includable: bool, reason: str = "",
                tracking: str | None = None) -> Assignment:
    return Assignment(group.ids, tracking or invoice.tracking, group.recipient, group.items, status, includable,
                      reason, invoice.carrier)


def _decide(group: _Group, picked: list[int], shared_by: list[int], invoices: list[Invoice],
            groups: list[_Group], candidates: list[list[int]], by_item: bool = False) -> Assignment:
    if len(picked) > 1:
        tracking = ", ".join(invoices[i].tracking for i in picked)
        return _assignment(group, invoices[picked[0]], STATUS_CHECK, False, "송장 여러 개", tracking)
    invoice = invoices[picked[0]]
    phone_ok, addr_ok = _mismatch(group, invoice)
    if len(shared_by) == 1:
        if phone_ok and addr_ok and by_item:
            return _assignment(group, invoice, STATUS_CHECK, True, REASON_BY_ITEM)
        if phone_ok and addr_ok:
            return _assignment(group, invoice, STATUS_MATCHED, True)
        return _assignment(group, invoice, STATUS_CHECK, True, _problem_reason(phone_ok, addr_ok))
    all_clean = all(len(candidates[g]) == 1 and all(_mismatch(groups[g], invoice))
                    for g in shared_by)
    if all_clean:
        return _assignment(group, invoice, STATUS_CHECK, True, "따로 주문했지만 한 상자로 보낸 것 같음")
    return _assignment(group, invoice, STATUS_CHECK, False, "여러 주문에 걸친 송장")


def match_invoices(orders: pd.DataFrame, invoices: list[Invoice]) -> MatchResult:
    """Pair each shipment group of the order table with the tracking number(s) of the same recipient."""
    no_tracking = tuple(i for i in invoices if not i.tracking)
    invoices = [i for i in invoices if i.tracking]
    groups = _group_orders(orders)
    candidates = [[i for i, inv in enumerate(invoices) if _is_candidate(g, inv)] for g in groups]
    by_item = _pick_by_item(groups, candidates, invoices)
    sharers: dict[int, list[int]] = defaultdict(list)
    for g, picked in enumerate(candidates):
        for i in picked:
            sharers[i].append(g)
    assignments, unshipped = [], []
    for g, group in enumerate(groups):
        if not candidates[g]:
            unshipped.append(UnshippedOrder(group.ids, group.recipient, group.items))
            continue
        shared = sharers[candidates[g][0]]
        assignments.append(_decide(group, candidates[g], shared, invoices, groups, candidates, g in by_item))
    leftovers = tuple(inv for i, inv in enumerate(invoices) if i not in sharers)
    return MatchResult(tuple(assignments), leftovers, tuple(unshipped), no_tracking)


# ---------------------------------------------------------------- upload workbook

def _upload_columns(columns: list[str]) -> list[str]:
    """All order columns; the three dispatch columns are added after the order number when absent."""
    result = list(columns)
    anchor = result.index(COL_ORDER_ID) + 1 if COL_ORDER_ID in result else (
        result.index(COL_ID) + 1 if COL_ID in result else 0)
    for name in (COL_METHOD, COL_CARRIER, COL_TRACKING):
        if name not in result:
            result.insert(anchor, name)
            anchor += 1
    return result


def _excel_value(value):
    if _blank(value):
        return None
    if isinstance(value, pd.Timestamp):
        return value.to_pydatetime()
    if hasattr(value, "item") and not isinstance(value, datetime):
        return value.item()
    return value


def build_upload_workbook(orders: pd.DataFrame, assignments, carrier: str = CARRIER_NAME,
                          method: str = DELIVERY_METHOD_LABEL) -> bytes:
    """xlsx for Naver's 엑셀 일괄발송: sheet 발송처리, header in row 1, only the assigned rows."""
    tracking_of = {pid: a.tracking for a in assignments for pid in a.product_order_ids}
    carrier_of_id = {pid: a.carrier for a in assignments for pid in a.product_order_ids}
    table = normalize_orders(orders)
    columns = _upload_columns(list(table.columns))
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = UPLOAD_SHEET
    sheet.append(columns)
    for row in table.to_dict("records"):
        tracking = tracking_of.get(_text(row.get(COL_ID)))
        if tracking is None:
            continue
        courier = carrier_of_id.get(_text(row.get(COL_ID))) or carrier
        row = {**row, COL_METHOD: method, COL_CARRIER: courier, COL_TRACKING: tracking}
        sheet.append([_excel_cell(row.get(c, ""), c) for c in columns])
    for position, name in enumerate(columns, start=1):
        if name in TEXT_COLUMNS:
            for row_number in range(2, sheet.max_row + 1):
                sheet.cell(row_number, position).number_format = "@"
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _excel_cell(value, column: str):
    if column in TEXT_COLUMNS or column in (COL_METHOD, COL_CARRIER):
        return _text(value)
    return _excel_value(value)
