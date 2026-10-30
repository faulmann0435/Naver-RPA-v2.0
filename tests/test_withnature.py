"""자연이랑 channel: export detection and conversion, courier codes, portal upload, purchase orders, saved orders."""
import json
from datetime import date, datetime
from io import BytesIO

import openpyxl
import pytest

from core.config_loader import load_config
from core.dictionary import DictionarySettings, ItemDictionary
from core.withnature import (
    build_portal_upload,
    courier_code,
    is_withnature_export,
    read_withnature_export,
    to_order_frame,
)
from tests.invoice_helpers import hanjin_row, hanjin_xlsx, order_export_xlsx, order_row
from tests.withnature_helpers import CRAB, OFFICE_MARK, SQUID, export_row, export_xlsx
from worker.naver_adapter import KST
from worker.order_run import Engine
from worker.paths import CONFIG_PATH
from worker.withnature_run import (
    load_saved_orders,
    run_withnature,
    save_orders,
    summary,
)

NOW = datetime(2026, 10, 30, 16, 30, tzinfo=KST)


# --- detection ------------------------------------------------------------------------------------

def test_detects_the_export_but_not_other_files():
    assert is_withnature_export(export_xlsx(export_row()))
    assert not is_withnature_export(hanjin_xlsx(hanjin_row(1, "111")))
    assert not is_withnature_export(order_export_xlsx(order_row("1")))
    assert not is_withnature_export(b"junk")


# --- conversion -----------------------------------------------------------------------------------

def _frame(*rows):
    return to_order_frame(read_withnature_export(export_xlsx(*rows)))


def test_columns_and_ids_follow_the_naver_export():
    frame = _frame(export_row(주문수량=2))
    row = frame.iloc[0]
    assert row["상품주문번호"] == row["주문번호"] == row["배송비 묶음번호"] == "O26100000001"
    assert row["상품번호"] == "자연이랑-2020100383"
    assert row["상품명"] == row["옵션정보"] == SQUID
    assert row["수량"] == 2 and frame["수량"].dtype == "int64"
    assert row["수취인명"] == "홍길동" and row["구매자명"] == "주문자" and row["주문상태"] == "발송대기"


def test_phone_falls_back_to_the_landline_and_zip_is_stripped():
    mobile = _frame(export_row(전화번호="02-000-0000"))
    assert (mobile.iloc[0]["수취인연락처1"], mobile.iloc[0]["수취인연락처2"]) == ("010-0000-0001", "02-000-0000")
    landline = _frame(export_row(휴대전화="", 전화번호="02-000-0000"))
    assert landline.iloc[0]["수취인연락처1"] == "02-000-0000"
    assert mobile.iloc[0]["우편번호"] == "13482"
    assert _frame(export_row(우편번호=1234)).iloc[0]["우편번호"] == "01234"


def test_stars_in_the_address_become_commas():
    address = _frame(export_row()).iloc[0]["통합배송지"]
    assert "(사아동, 자차아파트)" in address and "*" not in address
    assert _frame(export_row(받는사람주소="가*나")).iloc[0]["통합배송지"] == "가,나"


@pytest.mark.parametrize(("message", "memo", "expected"), [
    ("문 앞", "", "문 앞"), ("", "전달", "전달"), ("문 앞", "전달", "문 앞 / 전달"), ("같음", "같음", "같음"), ("", "", ""),
])
def test_message_and_memo_are_joined(message, memo, expected):
    assert _frame(export_row(배송메시지=message, 전달메모=memo)).iloc[0]["배송메세지"] == expected


def test_the_office_address_never_reaches_the_table():
    frame = _frame(export_row())
    assert OFFICE_MARK not in frame.to_csv() and "1588-0000" not in frame.to_csv()


def test_not_an_export_is_a_readable_error():
    with pytest.raises(ValueError, match="자연이랑"):
        read_withnature_export(hanjin_xlsx(hanjin_row(1, "111")))


# --- courier codes and the portal upload ----------------------------------------------------------

@pytest.mark.parametrize(("name", "code"), [
    ("한진택배", "03"), ("한진", "03"), ("CJ대한통운", "02"), ("씨제이대한통운", "02"), ("대한통운", "02"),
    ("로젠택배", "04"), ("롯데택배", "05"), ("우체국택배", "01"), ("우체국", "01"), ("DHL", "06"), ("FedEx", "07"),
    ("UPS", "08"), ("농협택배", "17"), ("경동택배", "12"), ("에이치피엠", "18"), ("한 진 택배", "03"),
    ("모르는택배", None), ("", None), (None, None), ("GROUPS", None),
])
def test_courier_code(name, code):
    assert courier_code(name) == code


def test_portal_upload_is_a_plain_text_sheet():
    data = build_portal_upload([("O26100000001", "000123456789", "03"), ("O26100000001", "999", "02")])
    rows = list(openpyxl.load_workbook(BytesIO(data)).active.iter_rows(values_only=True))
    assert rows == [("주문번호", "운송장번호", "택배업체코드"), ("O26100000001", "000123456789", "03"),
                    ("O26100000001", "999", "02")]


# --- purchase orders ------------------------------------------------------------------------------

def _engine():
    return Engine(load_config(str(CONFIG_PATH), password="1111"), ItemDictionary.empty(), DictionarySettings(), "test")


def _run(tmp_path, *rows, now=NOW):
    return run_withnature(export_xlsx(*rows), engine_loader=_engine, orders_file=tmp_path / "orders.json",
                          now=lambda: now)


def _sheet(file):
    return list(openpyxl.load_workbook(BytesIO(file["data"].getvalue())).active.iter_rows(values_only=True))


def test_two_products_make_two_prefixed_vendor_files(tmp_path):
    run = _run(tmp_path, export_row("O1", "김가나"), export_row("O2", "이다라", 상품코드="2020100263", 상품명=CRAB),
               export_row("O3", "박마바"))
    names = sorted(f["filename"] for f in run.files)
    assert len(names) == 2 and all(n.startswith("자연이랑_") for n in names)
    counts = dict(run.vendor_counts)
    assert sorted(counts.values()) == [1, 2] and run.order_count == 3
    squid = next(f for f in run.files if counts[f["filename"]] == 2)
    body = _sheet(squid)
    flat = "\n".join(str(c) for r in body for c in r)
    assert "김가나" in flat and "박마바" in flat and "이다라" not in flat
    assert "(사아동, 자차아파트)" in flat and "010-0000-0001" in flat
    assert "국내산 생물 갑오징어 1kg" in flat and "문 앞에 놓아 주세요" in flat and "최고다농수산" in flat
    assert OFFICE_MARK not in flat


def test_summary_counts_and_unmatched_warning(tmp_path):
    text = summary(_run(tmp_path, export_row("O1"), export_row("O2", 상품코드="2020100263", 상품명=CRAB)))
    assert "자연이랑 발주서 2개" in text and "주문 2건" in text and "미등록 품목 2건" in text and "앱에서 등록" in text


# --- saved orders ---------------------------------------------------------------------------------

def test_orders_are_saved_merged_and_expire_after_30_days(tmp_path):
    path = tmp_path / "orders.json"
    old = _frame(export_row("O1"))
    save_orders(old, date(2026, 9, 20), path)
    save_orders(_frame(export_row("O2")), date(2026, 10, 25), path)
    save_orders(_frame(export_row("O3")), NOW.date(), path)  # O1 is 40 days old now
    assert sorted(json.loads(path.read_text(encoding="utf-8"))) == ["O2", "O3"]
    loaded = load_saved_orders(path, NOW.date())
    assert sorted(loaded["상품주문번호"]) == ["O2", "O3"] and loaded["수량"].tolist() == [1, 1]
    assert load_saved_orders(path, date(2026, 12, 30)).empty


def test_repeated_order_numbers_are_still_included_and_named(tmp_path):
    first = _run(tmp_path, export_row("O1", "김가나"))
    assert first.repeated_names == () and "전에 받은 파일" not in summary(first)
    second = _run(tmp_path, export_row("O1", "김가나"), export_row("O2", "이다라"))
    assert second.order_count == 2 and second.repeated_names == ("김가나",)
    text = summary(second)
    assert "전에 받은 파일에도 있던 주문 1건: 김가나" in text and "두 번 보내지 않게" in text


def test_repeated_names_are_cut_at_ten(tmp_path):
    rows = [export_row(f"O{i}", f"고객{i}") for i in range(13)]
    _run(tmp_path, *rows)
    text = summary(_run(tmp_path, *rows))
    assert "외 3건" in text and "고객12" not in text
