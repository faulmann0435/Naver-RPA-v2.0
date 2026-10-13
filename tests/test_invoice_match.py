"""Invoice matching: Hanjin rows <-> Naver order rows, and the upload workbook."""
from io import BytesIO

import openpyxl
import pandas as pd
import pytest

from core.invoice_match import (
    STATUS_CHECK,
    STATUS_MATCHED,
    build_upload_workbook,
    match_invoices,
    normalize_orders,
    read_invoice_file,
    read_order_export,
)
from tests.invoice_helpers import (
    hanjin_row,
    hanjin_xlsx,
    order_export_xlsx,
    order_row,
    orders,
)


def _match(order_rows, hanjin_rows):
    return match_invoices(normalize_orders(orders(*order_rows)),
                          read_invoice_file(hanjin_xlsx(*hanjin_rows)))


def test_simple_match_ignores_hanjin_suffix_and_spacing():
    result = _match([order_row("1", "홍 길동", 통합배송지="서울특별시  가나구 다라로 1 (마바동)  101호")],
                    [hanjin_row(1, "111")])
    (a,) = result.assignments
    assert a.status == STATUS_MATCHED and a.tracking == "111" and a.product_order_ids == ("1",)
    assert result.invoices_without_order == () and result.orders_without_invoice == ()


def test_bundled_items_share_one_tracking():
    rows = [order_row("1", **{"배송비 묶음번호": "G"}), order_row("2", **{"배송비 묶음번호": "G"})]
    (a,) = _match(rows, [hanjin_row(1, "111")]).assignments
    assert a.product_order_ids == ("1", "2") and a.status == STATUS_MATCHED and "외 1건" in a.items


def test_group_without_bundle_number_falls_back_to_recipient():
    rows = [order_row("1", **{"배송비 묶음번호": ""}), order_row("2", **{"배송비 묶음번호": ""})]
    (a,) = _match(rows, [hanjin_row(1, "111")]).assignments
    assert a.product_order_ids == ("1", "2")


def test_safe_number_second_phone_matches():
    rows = [order_row("1", 수취인연락처1="0502-2764-8961")]
    result = _match(rows, [hanjin_row(1, "111", **{"받으시는 분 전화": "050227648961"})])
    assert result.assignments[0].status == STATUS_MATCHED
    rows = [order_row("1", 수취인연락처1="010-0000-0009", 수취인연락처2="010-0000-0001")]
    assert _match(rows, [hanjin_row(1, "111")]).assignments[0].status == STATUS_MATCHED


def test_phone_mismatch_needs_check_but_includable():
    (a,) = _match([order_row("1")], [hanjin_row(1, "111", **{"받으시는 분 전화": "01099999999"})]).assignments
    assert (a.status, a.includable, a.reason) == (STATUS_CHECK, True, "전화번호 다름")


def test_address_mismatch_needs_check_but_includable():
    (a,) = _match([order_row("1")], [hanjin_row(1, "111", 받는분총주소="전혀 다른 주소 9 [동, 1]")]).assignments
    assert (a.status, a.includable, a.reason) == (STATUS_CHECK, True, "주소 다름")


def test_two_boxes_for_one_group_are_not_includable():
    (a,) = _match([order_row("1")], [hanjin_row(1, "111"), hanjin_row(2, "222")]).assignments
    assert (a.status, a.includable, a.reason) == (STATUS_CHECK, False, "송장 여러 개")
    assert a.tracking == "111, 222"


def test_one_box_for_two_groups_with_different_phone_is_not_includable():
    rows = [order_row("1"), order_row("2", 수취인연락처1="010-0000-0005")]
    result = _match(rows, [hanjin_row(1, "111")])
    assert [(a.includable, a.reason) for a in result.assignments] == [(False, "여러 주문에 걸친 송장")] * 2


def test_one_box_for_two_clean_groups_is_includable_with_note():
    result = _match([order_row("1"), order_row("2")], [hanjin_row(1, "111")])
    assert all(a.includable and a.status == STATUS_CHECK and "한 상자" in a.reason for a in result.assignments)


def test_invoice_without_order_and_order_without_invoice():
    result = _match([order_row("1", "김철수")], [hanjin_row(1, "111", "이영희")])
    assert [i.name for i in result.invoices_without_order] == ["이영희"]
    assert [o.product_order_ids for o in result.orders_without_invoice] == [("1",)]
    assert result.assignments == ()


def test_zip_leading_zero_int_vs_text():
    table = normalize_orders(orders(order_row("1", 우편번호=8239)))
    assert table.loc[0, "우편번호"] == "08239"
    (a,) = match_invoices(table, read_invoice_file(hanjin_xlsx(hanjin_row(1, "111")))).assignments
    assert a.status == STATUS_MATCHED


def test_shipped_or_not_ready_orders_are_ignored():
    rows = [order_row("1", 송장번호="999"), order_row("2", 주문상태="배송중"), order_row("3", 수취인명="박민수")]
    result = _match(rows, [hanjin_row(1, "111", "박민수")])
    assert [a.product_order_ids for a in result.assignments] == [("3",)]


def test_hanjin_header_row_is_found_and_empty_tracking_skipped():
    data = hanjin_xlsx(hanjin_row(1, "111"), hanjin_row(2, None))
    assert [i.tracking for i in read_invoice_file(data)] == ["111"]
    with pytest.raises(ValueError, match="운송장번호"):
        read_invoice_file(hanjin_xlsx(hanjin_row(1, "1"), columns=["순번", "받으시는 분"]))
    with pytest.raises(ValueError, match="xlsx"):
        read_invoice_file(b"not an excel file")


def test_order_export_skips_instruction_row_and_keeps_text():
    data = order_export_xlsx(order_row("1", 우편번호=8239, 수취인연락처1="010-0000-0001"))
    table = read_order_export(data)
    assert list(table["상품주문번호"]) == ["1"] and table.loc[0, "우편번호"] == "08239"
    assert "안내" not in "".join(table.columns)


def _shrink_declared_size(data: bytes) -> bytes:
    """Rewrite the sheet's <dimension> to A1, as Naver's real export does."""
    import re
    import zipfile
    source, out = zipfile.ZipFile(BytesIO(data)), BytesIO()
    with zipfile.ZipFile(out, "w") as target:
        for item in source.infolist():
            body = source.read(item.filename)
            if item.filename.startswith("xl/worksheets/sheet"):
                body = re.sub(rb'<dimension ref="[^"]*"/>', b'<dimension ref="A1"/>', body)
            target.writestr(item, body)
    return out.getvalue()


def test_order_export_reads_past_a_wrong_declared_size():
    data = _shrink_declared_size(order_export_xlsx(order_row("1"), order_row("2")))
    assert list(read_order_export(data)["상품주문번호"]) == ["1", "2"]


def test_upload_workbook_shape():
    table = normalize_orders(orders(order_row("1", 수량=2), order_row("2", "김철수"), order_row("3")))
    result = match_invoices(table, read_invoice_file(hanjin_xlsx(hanjin_row(1, "123456789012"))))
    data = build_upload_workbook(table, result.includable())
    sheet = openpyxl.load_workbook(BytesIO(data)).active
    assert sheet.title == "발송처리"
    header = [c.value for c in sheet[1]]
    assert header == list(table.columns)  # all columns, original order
    rows = [dict(zip(header, [c.value for c in r], strict=True)) for r in sheet.iter_rows(min_row=2)]
    assert [r["상품주문번호"] for r in rows] == ["1", "3"]
    assert {(r["배송방법"], r["택배사"], r["송장번호"]) for r in rows} == {("택배,등기,소포", "한진택배", "123456789012")}
    assert rows[0]["우편번호"] == "08239" and isinstance(rows[0]["우편번호"], str)
    assert rows[0]["수량"] == 2
    assert all(c.data_type == "s" for c in sheet[2] if c.value not in (None, 2))


def test_upload_workbook_adds_dispatch_columns_for_api_table():
    api_table = pd.DataFrame({"상품주문번호": [1], "주문번호": [2], "우편번호": [8239], "수취인명": ["홍길동"]})
    assignment = _match([order_row("1")], [hanjin_row(1, "111")]).assignments[0]
    data = build_upload_workbook(api_table, [assignment])
    sheet = openpyxl.load_workbook(BytesIO(data)).active
    assert [c.value for c in sheet[1]] == ["상품주문번호", "주문번호", "배송방법", "택배사", "송장번호", "우편번호", "수취인명"]
    assert [c.value for c in sheet[2]][:5] == ["1", "2", "택배,등기,소포", "한진택배", "111"]
